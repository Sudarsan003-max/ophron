import os
import sys
import csv
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Colors
COLOR_NAVY = RGBColor(3, 33, 71)      # #032147 Primary Brand
COLOR_GOLD = RGBColor(183, 163, 139)  # #B7A38B Brand Accent / Champagne
COLOR_BODY = RGBColor(45, 55, 72)     # #2D3748 Readable Dark Slate
COLOR_MUTED = RGBColor(113, 128, 150) # #718096 Secondary Muted
HEX_NAVY = "032147"
HEX_GOLD = "B7A38B"
HEX_BG_LIGHT = "F4F6F9"
HEX_BORDER = "CBD5E0"
HEX_ZEBRA = "F8FAFC"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_left_border(cell, color_hex="032147", sz="36"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

def format_table(table, col_widths, headers, data, align_cols=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_row = table.rows[0]
    trPr = header_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for idx, name in enumerate(headers):
        cell = header_row.cells[idx]
        cell.text = name
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        if align_cols and idx in align_cols:
            p.alignment = align_cols[idx]
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
            run.font.name = "Arial"

    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg = HEX_ZEBRA if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            if align_cols and c_idx in align_cols:
                p.alignment = align_cols[c_idx]
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_BODY
                run.font.name = "Arial"

    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def add_callout(doc, title, text, border_color=HEX_NAVY):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, HEX_BG_LIGHT)
    set_cell_left_border(cell, color_hex=border_color, sz="32")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"✦ {title}\n")
    run_t.font.bold = True
    run_t.font.size = Pt(9.5)
    run_t.font.color.rgb = COLOR_NAVY if border_color == HEX_NAVY else COLOR_GOLD
    run_t.font.name = "Arial"
    run_b = p.add_run(text)
    run_b.font.size = Pt(9)
    run_b.font.color.rgb = COLOR_BODY
    run_b.font.name = "Arial"
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Helper functions initialized successfully.")
